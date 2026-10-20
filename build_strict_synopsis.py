"""
Build the strict 5-section News Shield Project Synopsis as specified in
'synopsis format.docx' and the user's requirements:

1. Title Page (with GNIOT NAAC A+ Logo header & Double Line Border)
2. Index & Notations / Nomenclature
3. 1. Introduction (Field, tech stack, problem statement, special terms)
4. 2. Literature Review (Shu et al., Perez-Rosas, Ahmed, Wang LIAR, CSI, FakeBERT, BERT, FakeNewsNet, Kaggle, Politifact)
5. 3. Objectives (8 bullet points)
6. 4. Methodology / Planning of Work (8 steps)
7. 5. References (IEEE standard based)
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_OUT = os.path.join(WORKSPACE_DIR, "News_Shield_Strict_Synopsis.docx")
PDF_OUT = os.path.join(WORKSPACE_DIR, "News_Shield_Project_Synopsis.pdf")
GNIOT_LOGO = os.path.join(WORKSPACE_DIR, "extracted_media", "word", "media", "image1.jpeg")


def set_cell_margins(cell, top=80, bottom=80, left=140, right=140):
    """Set cell internal padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_borders(table, color="B0BEC5", sz="4"):
    """Set clean thin borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_table_zero_indent(table):
    """Ensure table has zero negative indent so it stays strictly within margins."""
    tblPr = table._tbl.tblPr
    tblInd = parse_xml(f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>')
    tblPr.append(tblInd)


def add_double_page_borders(section):
    """Add GNIOT standard double-line page borders matching synopsis format.docx."""
    sectPr = section._sectPr
    borders_xml = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:left w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:bottom w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:right w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(borders_xml)


def generate_strict_synopsis():
    doc = docx.Document()

    # Page Margins: A4, Left 1.0 in (2.54 cm), Top 1.0 in (2.54 cm), Right 0.75 in (1.9 cm), Bottom 0.8 in (2.0 cm)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.8)

    add_double_page_borders(section)

    # Base Normal Style: Times New Roman, 12pt, 1.35 line spacing
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(3)

    # -------------------------------------------------------------
    # PAGE 1: TITLE PAGE
    # -------------------------------------------------------------
    t_header = doc.add_table(rows=1, cols=2)
    t_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_header.autofit = False
    set_table_zero_indent(t_header)

    c_logo, c_text = t_header.rows[0].cells
    c_logo.width = Inches(1.8)
    c_text.width = Inches(4.5)
    set_cell_margins(c_logo, top=40, bottom=40, left=40, right=40)
    set_cell_margins(c_text, top=40, bottom=40, left=40, right=40)

    # Insert GNIOT Logo
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(GNIOT_LOGO):
        p_logo.add_run().add_picture(GNIOT_LOGO, width=Inches(1.65))

    # College Name
    p_inst = c_text.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_after = Pt(0)
    r1 = p_inst.add_run("GREATER NOIDA INSTITUTE OF TECHNOLOGY\n")
    r1.bold = True
    r1.font.size = Pt(13.5)
    r1.font.name = 'Times New Roman'
    r2 = p_inst.add_run("(ENGINEERING INSTITUTE)\n")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.name = 'Times New Roman'
    r3 = p_inst.add_run("Plot No-7, Knowledge Park-II, Greater Noida")
    r3.font.size = Pt(10)
    r3.font.name = 'Times New Roman'

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(22)
    r_bar = p_div.add_run("―" * 44)
    r_bar.font.color.rgb = RGBColor(160, 160, 160)

    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.2
    p_title.paragraph_format.space_after = Pt(16)
    r_t = p_title.add_run("News Shield: An Intelligent System for Detection and Verifying Fake News")
    r_t.bold = True
    r_t.font.size = Pt(21)
    r_t.font.name = 'Times New Roman'
    r_t.font.color.rgb = RGBColor(15, 44, 89)

    # Document Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.line_spacing = 1.15
    p_sub.paragraph_format.space_after = Pt(14)
    r_s = p_sub.add_run("PROJECT SYNOPSIS\nOF MAJOR PROJECT")
    r_s.bold = True
    r_s.font.size = Pt(13)
    r_s.font.name = 'Times New Roman'

    # Degree Information
    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.line_spacing = 1.15
    p_deg.paragraph_format.space_after = Pt(14)
    r_d1 = p_deg.add_run("BACHELOR OF TECHNOLOGY\n")
    r_d1.bold = True
    r_d1.font.size = Pt(13.5)
    r_d1.font.name = 'Times New Roman'
    r_d2 = p_deg.add_run("COMPUTER SCIENCE & ENGINEERING\n")
    r_d2.bold = True
    r_d2.font.size = Pt(14.5)
    r_d2.font.name = 'Times New Roman'
    r_d3 = p_deg.add_run("Batch: 2022 - 2026")
    r_d3.bold = True
    r_d3.font.size = Pt(12)
    r_d3.font.name = 'Times New Roman'

    # Students Table (Width: 5.8 in, safe margin buffer)
    t_stud = doc.add_table(rows=1, cols=2)
    t_stud.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_stud.autofit = False
    set_table_zero_indent(t_stud)
    set_table_borders(t_stud, color="0F2C59", sz="6")

    h_s0, h_s1 = t_stud.rows[0].cells
    h_s0.width = Inches(3.4)
    h_s1.width = Inches(2.4)
    h_s0.text = "Name of Student"
    h_s1.text = "Class Roll Number"
    set_cell_shading(h_s0, "0F2C59")
    set_cell_shading(h_s1, "0F2C59")
    for c in [h_s0, h_s1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(c, top=80, bottom=80, left=140, right=140)

    students = [
        ("Aryan Chourasia", "2501320100070"),
        ("Ayush kumar singh", "2501320100080"),
        ("Arpita Dixit", "2501320100067"),
        ("Ananya Singh", "2501320100044")
    ]
    for name, roll in students:
        r_c0, r_c1 = t_stud.add_row().cells
        r_c0.width = Inches(3.4)
        r_c1.width = Inches(2.4)
        r_c0.text = name
        r_c1.text = roll
        r_c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in [r_c0, r_c1]:
            c.paragraphs[0].runs[0].font.size = Pt(11)
            c.paragraphs[0].runs[0].font.name = 'Times New Roman'
            set_cell_margins(c, top=50, bottom=50, left=140, right=140)

    # Supervisor / Guide
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(20)
    p_guide.paragraph_format.space_after = Pt(20)
    r_g = p_guide.add_run("SUPERVISOR / GUIDE: Mr. Rohit Vaid")
    r_g.bold = True
    r_g.font.size = Pt(12.5)
    r_g.font.name = 'Times New Roman'

    # Affiliated University & Year
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.line_spacing = 1.15
    p_univ.paragraph_format.space_after = Pt(0)
    r_u1 = p_univ.add_run("Greater Noida Institute of Technology, Greater Noida\n")
    r_u1.bold = True
    r_u1.font.size = Pt(13)
    r_u1.font.name = 'Times New Roman'
    r_u2 = p_univ.add_run("Dr. A.P.J. Abdul Kalam Technical University, Lucknow\n")
    r_u2.bold = True
    r_u2.font.size = Pt(14)
    r_u2.font.name = 'Times New Roman'
    r_u3 = p_univ.add_run("October 2026")
    r_u3.font.size = Pt(11.5)
    r_u3.font.name = 'Times New Roman'

    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 2: INDEX & NOTATIONS / NOMENCLATURE
    # -------------------------------------------------------------
    p_idx_h = doc.add_paragraph()
    p_idx_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_idx_h.paragraph_format.space_after = Pt(10)
    r_ih = p_idx_h.add_run("INDEX")
    r_ih.bold = True
    r_ih.font.size = Pt(14)
    r_ih.font.name = 'Times New Roman'

    # Width: 4.8 in + 1.2 in = 6.0 in (Cleanly inside 6.52 in text area)
    t_idx = doc.add_table(rows=1, cols=2)
    t_idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_idx.autofit = False
    set_table_zero_indent(t_idx)
    set_table_borders(t_idx, color="CBD5E1", sz="4")

    h0, h1 = t_idx.rows[0].cells
    h0.width = Inches(4.8)
    h1.width = Inches(1.2)
    h0.text = "Section Title"
    h1.text = "Page No."
    set_cell_shading(h0, "F1F5F9")
    set_cell_shading(h1, "F1F5F9")
    for c in [h0, h1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.name = 'Times New Roman'
        set_cell_margins(c, top=60, bottom=60, left=160, right=160)

    index_items = [
        ("1. Introduction", "3"),
        ("2. Literature Review", "4"),
        ("3. Objectives of the Project", "5"),
        ("4. Methodology / Planning of Work", "6"),
        ("5. References", "7")
    ]

    for title, pg in index_items:
        r0, r1 = t_idx.add_row().cells
        r0.width = Inches(4.8)
        r1.width = Inches(1.2)
        r0.text = title
        r1.text = pg
        r0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in [r0, r1]:
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            set_cell_margins(c, top=50, bottom=50, left=160, right=160)

    p_not_h = doc.add_paragraph()
    p_not_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_not_h.paragraph_format.space_before = Pt(24)
    p_not_h.paragraph_format.space_after = Pt(10)
    r_nh = p_not_h.add_run("NOTATIONS / NOMENCLATURE")
    r_nh.bold = True
    r_nh.font.size = Pt(13)
    r_nh.font.name = 'Times New Roman'

    t_not = doc.add_table(rows=1, cols=2)
    t_not.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_not.autofit = False
    set_table_zero_indent(t_not)
    set_table_borders(t_not, color="CBD5E1", sz="4")

    n0, n1 = t_not.rows[0].cells
    n0.width = Inches(1.8)
    n1.width = Inches(4.2)
    n0.text = "Abbreviation"
    n1.text = "Technical Definition"
    set_cell_shading(n0, "F1F5F9")
    set_cell_shading(n1, "F1F5F9")
    for c in [n0, n1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.name = 'Times New Roman'
        set_cell_margins(c, top=50, bottom=50, left=140, right=140)

    notations = [
        ("ML", "Machine Learning"),
        ("NLP", "Natural Language Processing"),
        ("TF-IDF", "Term Frequency - Inverse Document Frequency"),
        ("SVM", "Support Vector Machine"),
        ("LSTM", "Long Short-Term Memory"),
        ("BERT", "Bidirectional Encoder Representations from Transformers")
    ]

    for abbr, desc in notations:
        r0, r1 = t_not.add_row().cells
        r0.width = Inches(1.8)
        r1.width = Inches(4.2)
        r0.text = abbr
        r1.text = desc
        r0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for c in [r0, r1]:
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            set_cell_margins(c, top=45, bottom=45, left=140, right=140)

    doc.add_page_break()

    # -------------------------------------------------------------
    # HELPER FUNCTIONS FOR SECTIONS 1 TO 5
    # -------------------------------------------------------------
    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 44, 89)
        return p

    def add_sub_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.name = 'Times New Roman'
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2.5)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        return p

    def add_bullet_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        return p

    # -------------------------------------------------------------
    # 1. INTRODUCTION (Max 3 pages)
    # -------------------------------------------------------------
    add_section_heading("1. INTRODUCTION")

    add_sub_heading("Background and problem statement")
    add_body_p(
        "Digital communication allows news and claims to spread quickly through social media, websites, and messaging applications. Alongside legitimate reporting, fabricated, distorted, or misleading material may also be presented as news. Readers may have difficulty checking every item, while manual fact-checking takes time. The project addresses this need by proposing a quick screening tool for news-style text."
    )
    add_body_p(
        "The proposed News Shield system will use supervised machine learning to classify a news article, headline, or social-media post into dataset-defined 'real' or 'fake' categories. It will learn patterns from labelled examples. The model output is an estimate based on those examples; it does not independently verify evidence, confirm the source, or establish that a claim is true."
    )

    add_sub_heading("Proposed approach and technology")
    add_body_p(
        "The text-processing pipeline may include cleaning, tokenization, stop-word handling, and stemming or lemmatization. Candidate features include TF-IDF, Bag-of-Words, or word embeddings such as Word2Vec or GloVe. The proposed baseline comparison is Logistic Regression, Naive Bayes, Support Vector Machine, and Random Forest. If time and computing resources permit, LSTM/Bi-LSTM and BERT-based models may be explored as advanced comparisons."
    )
    add_body_p(
        "The project proposes an interactive web application in which a user can enter a headline or article and receive a 'Real' or 'Fake' prediction with a model confidence score. Flask or Django or Streamlit may be used for the application, subject to the team's implementation choice. Any score will be labelled as a model estimate, not a verified probability of truth. The system should help readers decide what to check; important claims still require original and independent sources."
    )

    add_sub_heading("Field of project and special technical terms")
    add_body_p(
        "The project belongs to the specialized field of Artificial Intelligence (AI) and Natural Language Processing (NLP). Key technical terms include:"
    )
    add_bullet_p(
        "• TF-IDF (Term Frequency-Inverse Document Frequency): A numerical statistic reflecting how important a word is to a document within a collection."
    )
    add_bullet_p(
        "• Porter Stemmer: A heuristic morphological algorithm removing word suffixes to map tokens to their base form."
    )
    add_bullet_p(
        "• Stylometry: The statistical measurement of writing style, including punctuation frequency, capitalization ratios, and sensationalist vocabulary."
    )

    add_sub_heading("Scope and expected contribution")
    add_body_p(
        "The initial scope is text classification using a documented labelled dataset. Social-context signals, source-credibility indicators, transformer models, and deployment are possible extensions and will be included only if they can be implemented and evaluated within the project scope. Expected deliverables are a documented data pipeline, a model comparison using held-out data, an interactive prototype, and a report of errors and limitations. No accuracy level is assumed in advance."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. LITERATURE REVIEW (Max 2 pages)
    # -------------------------------------------------------------
    add_section_heading("2. LITERATURE REVIEW")
    add_body_p(
        "Shu et al. survey fake-news detection from a data-mining perspective and organize prior work around news content, social context, datasets, and evaluation [1]. Their review highlights that text is only one possible signal. This supports the proposed text-focused baseline while motivating a clear statement that it does not use all available context."
    )
    add_body_p(
        "Pérez-Rosas et al. introduced datasets covering multiple online news domains and evaluated linguistic approaches for automatic detection [2]. Ahmed, Traoré, and Saad studied n-gram analysis with machine-learning classifiers [3]. These works motivate testing conventional text features and classifiers, but their results belong to their specific datasets and experimental settings; they do not guarantee the same performance for this project."
    )
    add_body_p(
        "Wang's LIAR benchmark contains manually labelled short political statements and provides a benchmark for fine-grained classification [4]. Ruchansky et al.'s CSI model combines article text with user-response and source-behaviour information [5]. These works show the difference between a text-only prototype and systems that incorporate context or social signals."
    )
    add_body_p(
        "Kaliyar et al. proposed FakeBERT, combining BERT representations with convolutional layers for social-media fake-news classification [6]. Devlin et al. introduced BERT as a bidirectional language representation model [8]. FakeNewsNet provides news content together with social and spatiotemporal information for research [7]. The team's supplied draft also identifies the Kaggle Fake and Real News dataset and PolitiFact as possible data or fact-checking resources [9], [10]. Availability, licensing, labels, and suitability must be checked before choosing a dataset."
    )
    add_body_p(
        "The reviewed work ranges from conventional n-gram models to contextual deep-learning and social-context methods. The proposed mini/major project will first establish a reproducible baseline, then compare selected alternatives if resources permit. Evaluation on held-out examples, error analysis, and human verification are necessary because a classifier can fail on new domains, satire, or changing language."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 3. OBJECTIVES (On its own dedicated page)
    # -------------------------------------------------------------
    add_section_heading("3. OBJECTIVES")
    add_body_p(
        "Based upon the literature review and the identified research problem, the primary goal of News Shield is to design, evaluate, and benchmark a reliable machine learning system for news credibility assessment. The general objective is broken down into eight logically connected parts:"
    )

    objectives = [
        ("1. Literature & State-of-the-Art Review:", "Study existing fake news detection approaches, benchmark datasets (LIAR, FakeNewsNet, ISOT), feature representations, and evaluation frameworks relevant to the selected project scope."),
        ("2. Dataset Selection & Curation:", "Select, balance, and document a labelled real/fake news dataset across diverse domains (politics, science, health, technology), recording provenance, license, and known limitations."),
        ("3. Text Preprocessing & Cleaning:", "Implement a standardized NLP pipeline including URL/HTML stripping, non-alphabetic filtering, stop-word elimination, and morphological stemming (Porter Stemmer)."),
        ("4. Feature Extraction & Representation:", "Compare TF-IDF (unigrams and bigrams with sublinear scaling) and Bag-of-Words representations to extract discriminating lexical and stylometric patterns."),
        ("5. Baseline Classifier Benchmarking:", "Train and evaluate multiple supervised classification models, including Multinomial Naive Bayes, Logistic Regression, Support Vector Machines (Linear SVM), and Random Forest."),
        ("6. Rigorous Empirical Evaluation:", "Measure generalization performance using 5-Fold Stratified Cross-Validation, test accuracy, precision, recall, F1-score, ROC-AUC, and confusion matrix analysis on held-out data."),
        ("7. Interactive Prototype Development:", "Build a responsive Streamlit web application providing real-time article verification, model confidence meters, stylometric indicators, and batch CSV processing."),
        ("8. Error Analysis & Limitation Reporting:", "Analyze false-positive and false-negative boundaries, model uncertainty, and domain-transfer constraints to ensure responsible deployment as a screening aid.")
    ]

    for title, desc in objectives:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        r_b = p.add_run(f"{title} ")
        r_b.bold = True
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(11)
        r_d = p.add_run(desc)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(11)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. METHODOLOGY / PLANNING OF WORK (On its own dedicated page)
    # -------------------------------------------------------------
    add_section_heading("4. METHODOLOGY / PLANNING OF WORK")
    add_body_p(
        "The project methodology comprises eight systematic development stages executed sequentially during project completion to ensure rigorous engineering and scientific reproducibility:"
    )

    steps = [
        ("Stage 1: Requirements Formulation & Scope Definition", "Define the problem boundaries, project scope, dataset permissions, evaluation criteria, and user interface requirements in consultation with the project guide."),
        ("Stage 2: Dataset Acquisition & Exploratory Analysis", "Acquire and verify a balanced corpus of real and fake news articles; inspect class distributions, missing values, text-length variations, and potential duplicate samples."),
        ("Stage 3: Data Partitioning & Preprocessing Pipeline", "Partition the dataset into stratified 80% training and 20% test subsets. Fit preprocessing transformations (URL stripping, tokenization, stop-word filtering, stemming) strictly on training data."),
        ("Stage 4: Feature Engineering & Vector Space Construction", "Construct high-dimensional numerical feature matrices using TF-IDF vectorization (unigrams + bigrams, 8,000 max features, sublinear TF scaling) and stylometric indicators."),
        ("Stage 5: Classifier Training & Hyperparameter Tuning", "Train candidate baseline classifiers (Multinomial Naive Bayes, Logistic Regression, Linear SVM with Platt probability calibration, Random Forest, and Ensemble Voting)."),
        ("Stage 6: Multi-Metric Model Validation & Benchmarking", "Evaluate models using 5-Fold Stratified Cross-Validation on training data and hold-out testing for Accuracy, Precision, Recall, F1-Score, and Confusion Matrix generation."),
        ("Stage 7: Web Application Interface Implementation", "Develop an interactive web dashboard (Streamlit) supporting real-time headline/article input, confidence meters, batch CSV analysis, and comparative benchmark visualizations."),
        ("Stage 8: Robustness Testing & Final Documentation", "Perform stress testing on challenging unseen news articles, conduct error boundary analysis, document operational limits, and compile project documentation.")
    ]

    for title, desc in steps:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        r_b = p.add_run(f"{title}: ")
        r_b.bold = True
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(11)
        r_d = p.add_run(desc)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(11)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 5. REFERENCES (IEEE standard based)
    # -------------------------------------------------------------
    add_section_heading("5. REFERENCES")

    references = [
        "[1] K. Shu, A. Sliva, S. Wang, J. Tang, and H. Liu, \"Fake News Detection on Social Media: A Data Mining Perspective,\" ACM SIGKDD Explorations Newsletter, vol. 19, no. 1, pp. 22-36, 2017, doi: 10.1145/3137597.3137600.",
        "[2] V. Pérez-Rosas, B. Kleinberg, A. Lefevre, and R. Mihalcea, \"Automatic Detection of Fake News,\" in Proc. 27th International Conference on Computational Linguistics (COLING), Santa Fe, NM, USA, 2018, pp. 3391-3401.",
        "[3] H. Ahmed, I. Traoré, and S. Saad, \"Detection of Online Fake News Using N-Gram Analysis and Machine Learning Techniques,\" in Intelligent, Secure, and Dependable Systems in Distributed and Cloud Environments, Springer, Cham, 2017, pp. 127-138, doi: 10.1007/978-3-319-69155-8_9.",
        "[4] W. Y. Wang, \"\"Liar, Liar Pants on Fire\": A New Benchmark Dataset for Fake News Detection,\" in Proc. 55th Annual Meeting of the Association for Computational Linguistics (ACL), Vancouver, Canada, 2017, pp. 422-426, doi: 10.18653/v1/P17-2067.",
        "[5] N. Ruchansky, S. Seo, and Y. Liu, \"CSI: A Hybrid Deep Model for Fake News Detection,\" in Proc. ACM Conference on Information and Knowledge Management (CIKM), Singapore, 2017, pp. 797-806, doi: 10.1145/3132847.3132877.",
        "[6] R. K. Kaliyar, A. Goswami, and P. Narang, \"FakeBERT: Fake News Detection in Social Media with a BERT-Based Deep Learning Approach,\" Multimedia Tools and Applications, vol. 80, no. 8, pp. 11765-11788, 2021, doi: 10.1007/s11042-020-10183-2.",
        "[7] K. Shu, D. Mahudeswaran, S. Wang, D. Lee, and H. Liu, \"FakeNewsNet: A Data Repository with News Content, Social Context, and Spatiotemporal Information for Studying Fake News on Social Media,\" Big Data, vol. 8, no. 3, pp. 171-188, 2020, doi: 10.1089/big.2020.0062.",
        "[8] J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, \"BERT: Pre-Training of Deep Bidirectional Transformers for Language Understanding,\" in Proc. NAACL-HLT, vol. 1, Minneapolis, MN, USA, 2019, pp. 4171-4186, doi: 10.18653/v1/N19-1423.",
        "[9] \"Fake and Real News Dataset,\" Kaggle. [Online]. Available: https://www.kaggle.com/datasets/clmentbisaillon/fake-and-real-news-dataset. Accessed: Oct. 16, 2026.",
        "[10] \"PolitiFact: Fact-Checking Journalism,\" [Online]. Available: https://www.politifact.com. Accessed: Oct. 16, 2026."
    ]

    for ref in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    doc.save(DOCX_OUT)
    print(f"[SUCCESS] Saved Strict DOCX: {DOCX_OUT}")


def convert_to_pdf():
    print("[INFO] Converting to PDF via Word COM...")
    cmd = [
        r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe",
        "-ExecutionPolicy", "Bypass",
        "-File", os.path.join(WORKSPACE_DIR, "convert_to_pdf.ps1"),
        "-InputPath", DOCX_OUT,
        "-OutputPath", PDF_OUT
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    print("PowerShell Output:", res.stdout)
    if res.stderr:
        print("PowerShell Error:", res.stderr)
    print(f"[SUCCESS] PDF generated at: {PDF_OUT}")


if __name__ == "__main__":
    generate_strict_synopsis()
    convert_to_pdf()
