"""
Build the complete, professionally formatted News Shield Major Project Synopsis
strictly following 'synopsis format.docx', incorporating:
- Double line page borders (thinThickSmallGap / thickThinSmallGap)
- GNIOT NAAC A+ Logo header (image1.jpeg)
- Fixed table margins and widths (no clipping or overflow)
- Complete academic sections: Introduction, Literature Review, Objectives, Category,
  DFD Level 0/1/2, System Architecture, Modules & Data Structures, H/W & S/W,
  Methodology, CPRC Endorsement (image3.jpeg), IEEE References.
- Conversion to News_Shield_Project_Synopsis.pdf via Word COM automation.
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_OUT = os.path.join(WORKSPACE_DIR, "News_Shield_Synopsis_Final.docx")
PDF_OUT = os.path.join(WORKSPACE_DIR, "News_Shield_Project_Synopsis.pdf")
GNIOT_LOGO = os.path.join(WORKSPACE_DIR, "extracted_media", "word", "media", "image1.jpeg")
CPRC_LOGO = os.path.join(WORKSPACE_DIR, "extracted_media", "word", "media", "image3.jpeg")


def set_cell_margins(cell, top=100, bottom=100, left=180, right=180):
    """Set cell internal padding in dxa (1 pt = 20 dxa)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Set standard thin borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_table_zero_indent(table):
    """Ensure table has zero negative indent so it stays strictly within margins."""
    tblPr = table._tbl.tblPr
    tblInd = parse_xml(f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>')
    tblPr.append(tblInd)


def add_page_borders(section):
    """Add GNIOT standard double line page borders matching synopsis format.docx."""
    sectPr = section._sectPr
    borders_xml = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:left w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:bottom w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'  <w:right w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(borders_xml)


def generate_document():
    doc = docx.Document()

    # Page Margins: A4, Left 1.0 in (2.54 cm), Top 1.0 in (2.54 cm), Right 0.75 in (1.9 cm), Bottom 0.8 in (2.0 cm)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.8)

    # Attach the standard double page border
    add_page_borders(section)

    # Base Normal Style: Times New Roman, 12pt, 1.25 line spacing
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(2.5)

    # -------------------------------------------------------------
    # 1. TITLE PAGE HEADER WITH GNIOT LOGO (Borderless Table)
    # -------------------------------------------------------------
    t_header = doc.add_table(rows=1, cols=2)
    t_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_header.autofit = False
    set_table_zero_indent(t_header)
    
    # Column widths: 1.8 in for logo, 4.5 in for college name
    c_logo, c_text = t_header.rows[0].cells
    c_logo.width = Inches(1.8)
    c_text.width = Inches(4.5)
    set_cell_margins(c_logo, top=40, bottom=40, left=40, right=40)
    set_cell_margins(c_text, top=40, bottom=40, left=40, right=40)

    # Add GNIOT NAAC A+ logo
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(GNIOT_LOGO):
        p_logo.add_run().add_picture(GNIOT_LOGO, width=Inches(1.65))

    # Add College details
    p_inst = c_text.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_after = Pt(0)
    
    r1 = p_inst.add_run("GREATER NOIDA INSTITUTE OF TECHNOLOGY\n")
    r1.bold = True
    r1.font.size = Pt(13.5)
    r1.font.name = 'Times New Roman'
    
    r2 = p_inst.add_run("(ENGINEERING INSTITUTE)\n")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.name = 'Times New Roman'
    
    r3 = p_inst.add_run("Plot No-7, Knowledge Park-II, Greater Noida")
    r3.font.size = Pt(10)
    r3.font.name = 'Times New Roman'

    # Divider line under header
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(20)
    r_bar = p_div.add_run("―" * 44)
    r_bar.font.color.rgb = RGBColor(160, 160, 160)

    # Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.2
    p_title.paragraph_format.space_after = Pt(16)
    r_t = p_title.add_run("News Shield: An Intelligent System for Detection and Verifying Fake News")
    r_t.bold = True
    r_t.font.size = Pt(21)
    r_t.font.name = 'Times New Roman'
    r_t.font.color.rgb = RGBColor(15, 44, 89)

    # Document Classification
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.line_spacing = 1.15
    p_sub.paragraph_format.space_after = Pt(12)
    r_s = p_sub.add_run("PROJECT SYNOPSIS\nOF MAJOR PROJECT")
    r_s.bold = True
    r_s.font.size = Pt(13)
    r_s.font.name = 'Times New Roman'

    # Degree Details
    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.line_spacing = 1.15
    p_deg.paragraph_format.space_after = Pt(18)
    r_d1 = p_deg.add_run("BACHELOR OF TECHNOLOGY\n")
    r_d1.bold = True
    r_d1.font.size = Pt(13.5)
    r_d1.font.name = 'Times New Roman'
    r_d2 = p_deg.add_run("COMPUTER SCIENCE & ENGINEERING")
    r_d2.bold = True
    r_d2.font.size = Pt(14.5)
    r_d2.font.name = 'Times New Roman'

    # Student Team Table (Width = 5.8 inches, centered, zero negative indent)
    t_stud = doc.add_table(rows=1, cols=2)
    t_stud.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_stud.autofit = False
    set_table_zero_indent(t_stud)
    set_table_borders(t_stud, color="0F2C59", sz="6")
    
    h_s0, h_s1 = t_stud.rows[0].cells
    h_s0.width = Inches(3.4)
    h_s1.width = Inches(2.4)
    h_s0.text = "Name of Student"
    h_s1.text = "University Roll Number"
    set_cell_shading(h_s0, "0F2C59")
    set_cell_shading(h_s1, "0F2C59")
    for c in [h_s0, h_s1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(c, top=80, bottom=80, left=140, right=140)

    students = [
        ("Aryan Chourasia", "2501320100070"),
        ("Ayush kumar singh", "2501320100080"),
        ("Arpita Dixit", "2501320100067"),
        ("Ananya Singh", "2501320100044")
    ]
    for name, roll in students:
        r_c0, r_c1 = t_stud.add_row().cells
        r_c0.width = Inches(3.4)
        r_c1.width = Inches(2.4)
        r_c0.text = name
        r_c1.text = roll
        r_c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in [r_c0, r_c1]:
            c.paragraphs[0].runs[0].font.size = Pt(11)
            c.paragraphs[0].runs[0].font.name = 'Times New Roman'
            set_cell_margins(c, top=50, bottom=50, left=140, right=140)

    # Supervisor / Guide
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(20)
    p_guide.paragraph_format.space_after = Pt(20)
    r_g = p_guide.add_run("SUPERVISOR / GUIDE: Mr. Rohit Vaid")
    r_g.bold = True
    r_g.font.size = Pt(12.5)
    r_g.font.name = 'Times New Roman'

    # Affiliated University & Year
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.line_spacing = 1.15
    p_univ.paragraph_format.space_after = Pt(0)
    r_u1 = p_univ.add_run("Greater Noida Institute of Technology, Greater Noida\n")
    r_u1.bold = True
    r_u1.font.size = Pt(13)
    r_u1.font.name = 'Times New Roman'
    r_u2 = p_univ.add_run("Dr. A.P.J. Abdul Kalam Technical University, Lucknow\n")
    r_u2.bold = True
    r_u2.font.size = Pt(14)
    r_u2.font.name = 'Times New Roman'
    r_u3 = p_univ.add_run("October 2026")
    r_u3.font.size = Pt(11.5)
    r_u3.font.name = 'Times New Roman'

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. INDEX & NOTATIONS (FIXED TABLE WIDTH & INDENT - NO OVERFLOW)
    # -------------------------------------------------------------
    p_idx_h = doc.add_paragraph()
    p_idx_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_idx_h.paragraph_format.space_after = Pt(10)
    r_ih = p_idx_h.add_run("INDEX")
    r_ih.bold = True
    r_ih.font.size = Pt(14)
    r_ih.font.name = 'Times New Roman'

    # Width: 5.0 in + 1.2 in = 6.2 in (Perfect fit inside 6.52 in text area)
    t_idx = doc.add_table(rows=1, cols=2)
    t_idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_idx.autofit = False
    set_table_zero_indent(t_idx)
    set_table_borders(t_idx, color="CBD5E1", sz="4")

    h0, h1 = t_idx.rows[0].cells
    h0.width = Inches(5.0)
    h1.width = Inches(1.2)
    h0.text = "Section Title"
    h1.text = "Page No."
    set_cell_shading(h0, "F1F5F9")
    set_cell_shading(h1, "F1F5F9")
    for c in [h0, h1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.name = 'Times New Roman'
        set_cell_margins(c, top=60, bottom=60, left=160, right=160)

    index_items = [
        ("1. Introduction & Problem Statement", "3"),
        ("2. Literature Review", "4"),
        ("3. Objectives of the Project", "5"),
        ("4. Project Category", "5"),
        ("5. System Analysis & Architecture Design (DFDs, Use Cases)", "6"),
        ("6. Complete Structure & Module Descriptions", "7"),
        ("7. Hardware & Software Requirements", "8"),
        ("8. Methodology & Planning of Work", "9"),
        ("9. Industry / Client Information", "10"),
        ("10. Future Scope & Enhancements", "10"),
        ("11. Central Project & Research Committee (CPRC) Endorsement", "11"),
        ("12. References (IEEE Format)", "11")
    ]

    for title, pg in index_items:
        r0, r1 = t_idx.add_row().cells
        r0.width = Inches(5.0)
        r1.width = Inches(1.2)
        r0.text = title
        r1.text = pg
        r0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in [r0, r1]:
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(11)
            set_cell_margins(c, top=45, bottom=45, left=160, right=160)

    # Notations Table
    p_not_h = doc.add_paragraph()
    p_not_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_not_h.paragraph_format.space_before = Pt(16)
    p_not_h.paragraph_format.space_after = Pt(8)
    r_nh = p_not_h.add_run("NOTATIONS / NOMENCLATURE")
    r_nh.bold = True
    r_nh.font.size = Pt(13)
    r_nh.font.name = 'Times New Roman'

    t_not = doc.add_table(rows=1, cols=2)
    t_not.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_not.autofit = False
    set_table_zero_indent(t_not)
    set_table_borders(t_not, color="CBD5E1", sz="4")

    n0, n1 = t_not.rows[0].cells
    n0.width = Inches(1.8)
    n1.width = Inches(4.4)
    n0.text = "Abbreviation"
    n1.text = "Technical Definition"
    set_cell_shading(n0, "F1F5F9")
    set_cell_shading(n1, "F1F5F9")
    for c in [n0, n1]:
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.name = 'Times New Roman'
        set_cell_margins(c, top=50, bottom=50, left=140, right=140)

    notations = [
        ("ML", "Machine Learning (Supervised Binary Text Classification)"),
        ("NLP", "Natural Language Processing (Tokenization, Stopwords, Stemming)"),
        ("TF-IDF", "Term Frequency - Inverse Document Frequency"),
        ("SVM", "Support Vector Machine (Linear Margin Classifier)"),
        ("LR", "Logistic Regression (L2-Regularized Sigmoidal Classifier)"),
        ("PAC", "Passive Aggressive Classifier (Online Streaming Algorithm)"),
        ("DFD", "Data Flow Diagram (Hierarchical Functional Decomposition)"),
        ("BERT", "Bidirectional Encoder Representations from Transformers"),
        ("ROC-AUC", "Receiver Operating Characteristic - Area Under Curve"),
        ("API / UI", "Application Programming Interface / Streamlit Web Interface")
    ]

    for abbr, desc in notations:
        r0, r1 = t_not.add_row().cells
        r0.width = Inches(1.8)
        r1.width = Inches(4.4)
        r0.text = abbr
        r1.text = desc
        r0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for c in [r0, r1]:
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10.5)
            set_cell_margins(c, top=35, bottom=35, left=140, right=140)

    doc.add_page_break()

    # -------------------------------------------------------------
    # HELPER FORMATTING FUNCTIONS
    # -------------------------------------------------------------
    def add_section_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 44, 89)
        return p

    def add_sub_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.name = 'Times New Roman'
        return p

    def add_para(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2.5)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        return p

    # -------------------------------------------------------------
    # 1. INTRODUCTION
    # -------------------------------------------------------------
    add_section_title("1. INTRODUCTION")

    add_sub_title("1.1 Background and Problem Statement")
    add_para(
        "Digital communication and modern social networking platforms enable news, breaking reports, and public claims to proliferate rapidly across the globe. Alongside legitimate and rigorous investigative journalism, fabricated, distorted, or sensationalist material is routinely manufactured and disseminated as genuine reporting. Everyday readers face formidable cognitive challenges when attempting to manually authenticate every article encountered online, while exhaustive manual fact-checking by human experts requires substantial investigation time and resources. The News Shield project addresses this pressing societal vulnerability by proposing an automated, rapid screening tool designed specifically for news-style text."
    )

    add_sub_title("1.2 Proposed Approach and Technology Stack")
    add_para(
        "The proposed News Shield system implements a supervised Machine Learning (ML) and Natural Language Processing (NLP) framework to classify incoming news articles, headlines, and social media posts into dataset-defined 'Real' (Authentic) or 'Fake' (Misinformation) categories. The end-to-end processing pipeline comprises automated web cleansing, tokenization, stopword elimination, and morphological stemming via the Porter Stemmer. For numeric representation, the system extracts high-dimensional unigram and bigram Term Frequency - Inverse Document Frequency (TF-IDF) feature vectors with sublinear scaling. The classification backend benchmarks baseline models including Logistic Regression, Multinomial Naive Bayes, Support Vector Machines (SVM), and an ensemble Soft Voting classifier. Predictions are surfaced through a modern, responsive web application (Streamlit) featuring real-time probabilistic confidence estimates."
    )

    add_sub_title("1.3 Scope, Operating Philosophy, and Expected Contributions")
    add_para(
        "The primary operational scope centers on text-based pattern recognition across verified corpora. Crucially, the system functions as an intelligent screening aid rather than an infallible, automated authority on absolute ground truth; critical claims and high-stakes reports continue to necessitate primary sourcing. Expected project deliverables comprise a reproducible data generation and cleaning pipeline, a rigorous multi-model empirical comparison using held-out test data and 5-fold cross-validation, an interactive web prototype, and a comprehensive analysis of errors and generalization boundaries across diverse domains."
    )

    # -------------------------------------------------------------
    # 2. LITERATURE REVIEW
    # -------------------------------------------------------------
    add_section_title("2. LITERATURE REVIEW")
    add_para(
        "The identification of deceptive text has emerged as a vibrant domain in data mining and computational linguistics. A thorough review of prior state-of-the-art literature was conducted to contextualize the proposed News Shield architecture:"
    )

    lit_data = [
        ("Shu et al. (2017) [1]:", "Surveyed fake news detection on social media from a data mining perspective, organizing techniques around news content, social context, and user engagement. Their analysis demonstrated that while social graph propagation offers valuable signals, content-level linguistic and lexical features form the fundamental, reproducible baseline for automated detection systems."),
        ("Pérez-Rosas et al. (2018) [2]:", "Introduced multi-domain cross-genre datasets and evaluated lexical, syntactic, and psycholinguistic markers (sentiment polarity, readability indices, and punctuation distributions), demonstrating that stylometric cues provide strong signals for separating authentic reporting from fabricated rumors."),
        ("Ahmed, Traoré, and Saad (2017) [3]:", "Investigated n-gram analysis combined with linear and nonlinear classifiers. Their experimental results proved that extracting unigrams and bigrams via TF-IDF achieves superior discriminative capability for distinguishing deceptive journalism from real news wire reports."),
        ("Wang (2017) [4]:", "Published the LIAR benchmark dataset consisting of 12,800 manually labeled short political claims from PolitiFact, establishing standard protocols for fine-grained veracity classification."),
        ("Ruchansky, Seo, and Liu (2017) [5]:", "Proposed the CSI hybrid deep model integrating textual features, user response sequences, and publisher behavioral patterns, highlighting the complementary nature of content and context."),
        ("Kaliyar, Goswami, and Narang (2021) [6]:", "Introduced FakeBERT, demonstrating that combining contextual bidirectional representations from transformers with convolutional neural network layers achieves high precision in detecting complex social media misinformation."),
        ("Devlin et al. (2019) [8] & FakeNewsNet [7]:", "Provided foundational insights into pre-trained transformer embeddings and comprehensive multi-attribute data repositories incorporating spatial-temporal attributes."),
        ("Kaggle & PolitiFact Resources [9, 10]:", "Documented standard public benchmarks of authentic news wire feeds (Reuters, AP) alongside verified political hoaxes, forming the structural foundation for balanced corpus generation.")
    ]

    for title, desc in lit_data:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run(f"• {title} ")
        r_b.bold = True
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(11)
        r_d = p.add_run(desc)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(11)

    # -------------------------------------------------------------
    # 3. OBJECTIVES OF THE PROJECT
    # -------------------------------------------------------------
    add_section_title("3. OBJECTIVES OF THE PROJECT")
    add_para("The specific, measurable objectives established for the News Shield Major Project are:")

    objectives = [
        "1. Study existing fake-news detection paradigms, feature extraction techniques, and benchmark datasets relevant to the computational scope.",
        "2. Curate and document a balanced, labelled real/fake news dataset (1,200 records across 8 distinct thematic domains) detailing provenance, label definitions, and domain scope.",
        "3. Implement an NLP text-processing pipeline incorporating regex filtering, stopword elimination, and Porter morphological stemming; extract TF-IDF unigram and bigram feature representations.",
        "4. Train, tune, and evaluate baseline classifiers: Logistic Regression, Multinomial Naive Bayes, and Support Vector Machines (SVM).",
        "5. Develop and evaluate an advanced Soft Voting Ensemble model that pools posterior probabilities across complementary linear and probabilistic architectures.",
        "6. Benchmark all candidate classifiers using Accuracy, Precision, Recall, F1-Score, ROC-AUC, 5-Fold Stratified Cross-Validation, and Confusion Matrix analysis on held-out test data.",
        "7. Build an interactive, user-friendly web interface (Streamlit) featuring single-click preset verification, live probability meters, stylometric marker breakdowns, and bulk CSV processing.",
        "8. Perform comprehensive error analysis on false positives and false negatives to delineate model generalization limits across unseen topics, writing styles, and sensationalism levels."
    ]

    for obj in objectives:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(obj)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    # -------------------------------------------------------------
    # 4. PROJECT CATEGORY
    # -------------------------------------------------------------
    add_section_title("4. PROJECT CATEGORY")
    add_para(
        "• Primary Technical Discipline: Artificial Intelligence & Machine Learning (Supervised Text Classification)\n"
        "• Sub-Discipline: Natural Language Processing (NLP), Stylometry & Information Retrieval\n"
        "• Application Domain: Web-Based Decision Support Systems & Digital Information Verification"
    )

    # -------------------------------------------------------------
    # 5. SYSTEM ANALYSIS & ARCHITECTURE DESIGN
    # -------------------------------------------------------------
    add_section_title("5. SYSTEM ANALYSIS & ARCHITECTURE DESIGN")

    add_sub_title("5.1 Data Flow Diagrams (DFDs)")
    add_para(
        "• Level 0 DFD (Context Level): The external user inputs raw news headlines or body paragraphs to the News Shield System. The system executes the vectorized inference pipeline and returns an authenticity badge ('Authentic' or 'Misinformation') accompanied by a calibrated confidence score and stylometric marker report.\n"
        "• Level 1 DFD (Subsystem Decomposition): Comprises four sequential processes: [P1.0 Data Ingestion & Sanitization] cleans URLs, HTML, and punctuation; [P2.0 NLP Feature Transformation] applies Porter Stemming and TF-IDF projection (8,000 features); [P3.0 Multi-Model Inference Engine] invokes the active classifier (Logistic Regression, Linear SVM, Ensemble) to compute decision probabilities; [P4.0 Visualization Subsystem] renders the UI badges, confidence meters, and CSV batch tables.\n"
        "• Level 2 DFD (Feature Engineering & Inference Engine): Transformed vectors are matched against calibrated model weights, Platt-scaled sigmoid probabilities are derived, and decision logic evaluates whether P(Real) exceeds the configured decision threshold (default: 0.50)."
    )

    add_sub_title("5.2 Use Case Analysis")
    add_para(
        "• Actor: End-User (Reader, Student, Journalist, or Content Moderator)\n"
        "• System Boundary: News Shield Web Interface & Inference API\n"
        "• Core Use Cases: UC-1 (Verify Single Article), UC-2 (Inspect Stylometric Markers), UC-3 (Select Active Classifier), UC-4 (Batch CSV News Classification), UC-5 (Review Benchmark & Viva Notes)."
    )

    add_sub_title("5.3 4-Layer System Architecture")
    add_para(
        "1. Presentation Layer: Built with Streamlit, providing interactive web dashboards, Plotly charts, and responsive CSS styling.\n"
        "2. Service / Preprocessing Layer: Python modules executing text cleaning, NLTK stopword pruning, and linguistic feature extraction.\n"
        "3. Machine Learning Core Layer: Scikit-Learn TF-IDF vectorizer (8,000 features) coupled with serialized model estimators (.pkl).\n"
        "4. Data & Serialization Layer: File-based persistence storing training datasets (sample_news_dataset.csv), serialized models (model.pkl, all_models.pkl), and evaluation metrics (model_metrics.json)."
    )

    # -------------------------------------------------------------
    # 6. COMPLETE STRUCTURE & MODULE DESCRIPTIONS
    # -------------------------------------------------------------
    add_section_title("6. COMPLETE STRUCTURE & MODULE DESCRIPTIONS")

    modules = [
        ("Module 1: Text Preprocessing & Cleaning (src/preprocessor.py)",
         "Responsible for text normalization and sanitization. Strips web hyperlinks, HTML markup tags, and non-alphanumeric punctuation. Converts characters to lowercase, tokenizes sentences, filters stopwords using the NLTK English corpus, and reduces terms to their root morphological stem via the Porter Stemmer. Also computes stylometric markers including word counts, uppercase character ratios, exclamation mark frequencies, and sensational keyword hits."),
        ("Module 2: Feature Engineering & Vectorization (models/vectorizer.pkl)",
         "Extracts TF-IDF feature representations across unigrams and bigrams (ngram_range=(1, 2)) capped at the top 8,000 most informative features. Sublinear term-frequency scaling is enabled (1 + log(tf)) to reduce distortion from hyper-frequent terms, with document frequency boundaries (min_df=2, max_df=0.98) to filter isolated typographical errors."),
        ("Module 3: Multi-Model Training & Benchmarking (train.py, src/evaluate.py)",
         "Implements 6 candidate machine learning classifiers: (i) L2-Regularized Logistic Regression, (ii) Calibrated Linear Support Vector Machine (LinearSVC with 5-fold Platt scaling), (iii) Soft Voting Ensemble combining LR, SVM, and Naive Bayes, (iv) Multinomial Naive Bayes, (v) Passive Aggressive Classifier, and (vi) Random Forest. Conducts 5-Fold Stratified Cross-Validation on training data and computes Precision, Recall, F1-Score, and ROC-AUC on held-out test data."),
        ("Module 4: Interactive Web Application & Batch Analyzer (app.py)",
         "Provides an end-to-end graphical web interface with 4 dedicated functional tabs: Real-Time News Verifier, Multi-Model Benchmark & Analytics Dashboard, Batch CSV Classifier, and Project Synopsis & Viva Notes. Features interactive model selection and adjustable decision threshold sliders.")
    ]

    for m_name, m_desc in modules:
        add_sub_title(m_name)
        add_para(m_desc)

    add_sub_title("6.1 Data Structures Employed")
    add_para(
        "• SciPy Compressed Sparse Row (CSR) Matrix: Encodes high-dimensional sparse TF-IDF text features (1,200 × 8,000) with minimal memory footprint.\n"
        "• Pandas DataFrame: Manages structured tabular article records, category tags, batch CSV inputs, and evaluation metrics.\n"
        "• NumPy ndarrays: Stores calibrated probability distributions, confusion matrix arrays, and decision function values.\n"
        "• Python Dictionaries & JSON: Serializes candidate model collections (all_models.pkl) and benchmark evaluation summaries (model_metrics.json)."
    )

    add_sub_title("6.2 Testing Strategies & Validation Process")
    add_para(
        "• 5-Fold Stratified Cross-Validation: Evaluates model generalization stability across 5 disjoint subsets preserving exact class balance.\n"
        "• 20% Held-Out Test Set: Measures final model generalization on 240 unseen articles, generating confusion matrices and classification reports.\n"
        "• 16-Article Real-World Generalization Benchmark: Rigorous testing against authentic, challenging real-world news stories (NASA, WHO, Federal Reserve, Supreme Court, 5G hoaxes, cancer cure scams) to ensure zero overfitting."
    )

    # -------------------------------------------------------------
    # 7. HARDWARE & SOFTWARE REQUIREMENTS
    # -------------------------------------------------------------
    add_section_title("7. HARDWARE & SOFTWARE REQUIREMENTS")

    add_sub_title("7.1 Software Requirements")
    add_para(
        "• Operating System: Windows 10/11, macOS, or Linux (Ubuntu 20.04+)\n"
        "• Programming Language: Python 3.9 or higher (Tested on Python 3.14)\n"
        "• Machine Learning Framework: Scikit-Learn 1.9+, SciPy 1.18+\n"
        "• Natural Language Toolkit: NLTK 3.10+ (PorterStemmer, English stopwords corpus)\n"
        "• Web Application Framework: Streamlit 1.61+\n"
        "• Data Handling & Visualization: Pandas 3.0+, NumPy 2.5+, Plotly 6.9+, Matplotlib 3.11+\n"
        "• Model Serialization: Joblib 1.5+"
    )

    add_sub_title("7.2 Hardware Specifications")
    add_para(
        "• Processor: Intel Core i3 / AMD Ryzen 3 or higher (Intel Core i5 recommended)\n"
        "• Primary Memory (RAM): Minimum 4 GB (8 GB recommended for batch vectorization)\n"
        "• Secondary Storage: 500 MB free hard disk space for code, virtual environment, and serialized models\n"
        "• Network: Standard internet connection for initial dependency installation"
    )

    # -------------------------------------------------------------
    # 8. METHODOLOGY & PLANNING OF WORK
    # -------------------------------------------------------------
    add_section_title("8. METHODOLOGY & PLANNING OF WORK")
    add_para("The project execution follows a structured milestone schedule aligned with academic major project rubrics:")

    phases = [
        ("Phase 1: Problem Definition & Literature Survey (Weeks 1-2)", "Review state-of-the-art NLP fake news literature; define project scope, baseline models, and evaluation criteria."),
        ("Phase 2: Dataset Selection & Synthesis (Weeks 3-4)", "Compile balanced 1,200-sample corpus across 8 journalistic and misinformation domains; perform EDA and class balance checks."),
        ("Phase 3: NLP Preprocessing & Feature Engineering (Weeks 5-6)", "Implement regex sanitization, stopword pruning, and Porter stemming; optimize 8,000-feature TF-IDF vectorizer."),
        ("Phase 4: Model Training & Hyperparameter Tuning (Weeks 7-8)", "Train Logistic Regression, Linear SVM, Multinomial Naive Bayes, Passive Aggressive, and Soft Voting Ensemble; conduct 5-Fold Stratified Cross-Validation."),
        ("Phase 5: Prototype Development (Weeks 9-10)", "Construct interactive Streamlit web application with model selector, probability meters, stylometric visualizers, and batch CSV processor."),
        ("Phase 6: Empirical Testing & Error Analysis (Weeks 11-12)", "Evaluate models on 16-article unseen real-world benchmark suite; analyze false-positive boundaries; finalize project synopsis and documentation.")
    ]

    for p_title, p_desc in phases:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        r_pt = p.add_run(f"• {p_title}: ")
        r_pt.bold = True
        r_pt.font.name = 'Times New Roman'
        r_pt.font.size = Pt(11)
        r_pd = p.add_run(p_desc)
        r_pd.font.name = 'Times New Roman'
        r_pd.font.size = Pt(11)

    # -------------------------------------------------------------
    # 9. INDUSTRY / CLIENT INFORMATION
    # -------------------------------------------------------------
    add_section_title("9. INDUSTRY / CLIENT INFORMATION")
    add_para(
        "• Project Type: In-House Academic Major Project (B.Tech Computer Science & Engineering)\n"
        "• Sponsoring Industry / Client: None (Self-initiated open research prototype)\n"
        "• Institutional Affiliation: Greater Noida Institute of Technology (GNIOT), Knowledge Park-II, Greater Noida"
    )

    # -------------------------------------------------------------
    # 10. FUTURE SCOPE & FURTHER ENHANCEMENT
    # -------------------------------------------------------------
    add_section_title("10. FUTURE SCOPE & FURTHER ENHANCEMENT")
    add_para(
        "1. Deep Transformer Representations: Fine-tune pre-trained contextual language models (BERT, RoBERTa, DistilBERT) to capture long-range contextual nuance and subtle semantic irony.\n"
        "2. Multimodal Verification: Expand detection to analyze accompanying multimedia (photographs, synthetic deepfake imagery, and video clips) alongside article text.\n"
        "3. Live Fact-Checking API Integration: Connect to external automated claim-checking services (Google Fact Check Tools API, ClaimReview) to retrieve authoritative source verdicts in real time.\n"
        "4. Multilingual Misinformation Detection: Extend tokenization and embedding architectures to support Indian regional languages (Hindi, Bengali, Tamil) to combat localized digital disinformation."
    )

    # -------------------------------------------------------------
    # 11. CENTRAL PROJECT & RESEARCH COMMITTEE (CPRC) ENDORSEMENT
    # -------------------------------------------------------------
    add_section_title("11. CENTRAL PROJECT & RESEARCH COMMITTEE (CPRC) ENDORSEMENT")
    
    t_cprc = doc.add_table(rows=1, cols=2)
    t_cprc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cprc.autofit = False
    set_table_zero_indent(t_cprc)
    
    c_cprc_logo, c_cprc_txt = t_cprc.rows[0].cells
    c_cprc_logo.width = Inches(1.5)
    c_cprc_txt.width = Inches(4.7)
    set_cell_margins(c_cprc_logo, top=40, bottom=40, left=40, right=40)
    set_cell_margins(c_cprc_txt, top=40, bottom=40, left=40, right=40)

    p_cp_logo = c_cprc_logo.paragraphs[0]
    p_cp_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(CPRC_LOGO):
        p_cp_logo.add_run().add_picture(CPRC_LOGO, width=Inches(1.2))

    p_cp_txt = c_cprc_txt.paragraphs[0]
    p_cp_txt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cp_txt.paragraph_format.line_spacing = 1.15
    r_cp1 = p_cp_txt.add_run("Central Project and Research Committee (CPRC)\n")
    r_cp1.bold = True
    r_cp1.font.size = Pt(11)
    r_cp1.font.name = 'Times New Roman'
    r_cp2 = p_cp_txt.add_run("Greater Noida Institute of Technology (Engineering Institute), Greater Noida GNIOT (132)\n")
    r_cp2.font.size = Pt(10)
    r_cp2.font.name = 'Times New Roman'
    r_cp3 = p_cp_txt.add_run("Dr. Shivani Dubey (Head CPRC)  |  Dr. Nancy Agrawal (CPRC Coordinator - CSE/IT/AI-ML)\nProf. (Dr.) Dhiraj Gupta (Director)")
    r_cp3.font.size = Pt(9.5)
    r_cp3.font.name = 'Times New Roman'

    # -------------------------------------------------------------
    # 12. REFERENCES (IEEE FORMAT)
    # -------------------------------------------------------------
    add_section_title("12. REFERENCES")

    references = [
        "[1] K. Shu, A. Sliva, S. Wang, J. Tang, and H. Liu, 'Fake News Detection on Social Media: A Data Mining Perspective,' ACM SIGKDD Explorations Newsletter, vol. 19, no. 1, pp. 22-36, 2017. doi: 10.1145/3137597.3137600.",
        "[2] V. Pérez-Rosas, B. Kleinberg, A. Lefevre, and R. Mihalcea, 'Automatic Detection of Fake News,' in Proc. 27th International Conference on Computational Linguistics (COLING), Santa Fe, NM, USA, 2018, pp. 3391-3401.",
        "[3] H. Ahmed, I. Traoré, and S. Saad, 'Detection of Online Fake News Using N-Gram Analysis and Machine Learning Techniques,' in Intelligent, Secure, and Dependable Systems in Distributed and Cloud Environments, Springer, Cham, 2017, pp. 127-138. doi: 10.1007/978-3-319-69155-8_9.",
        "[4] W. Y. Wang, ''Liar, Liar Pants on Fire': A New Benchmark Dataset for Fake News Detection,' in Proc. 55th Annual Meeting of the Association for Computational Linguistics (ACL), Vancouver, Canada, 2017, pp. 422-426. doi: 10.18653/v1/P17-2067.",
        "[5] N. Ruchansky, S. Seo, and Y. Liu, 'CSI: A Hybrid Deep Model for Fake News Detection,' in Proc. ACM Conference on Information and Knowledge Management (CIKM), Singapore, 2017, pp. 797-806. doi: 10.1145/3132847.3132877.",
        "[6] R. K. Kaliyar, A. Goswami, and P. Narang, 'FakeBERT: Fake News Detection in Social Media with a BERT-Based Deep Learning Approach,' Multimedia Tools and Applications, vol. 80, no. 8, pp. 11765-11788, 2021. doi: 10.1007/s11042-020-10183-2.",
        "[7] K. Shu, D. Mahudeswaran, S. Wang, D. Lee, and H. Liu, 'FakeNewsNet: A Data Repository with News Content, Social Context, and Spatiotemporal Information for Studying Fake News on Social Media,' Big Data, vol. 8, no. 3, pp. 171-188, 2020. doi: 10.1089/big.2020.0062.",
        "[8] J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, 'BERT: Pre-Training of Deep Bidirectional Transformers for Language Understanding,' in Proc. NAACL-HLT, vol. 1, Minneapolis, MN, USA, 2019, pp. 4171-4186. doi: 10.18653/v1/N19-1423.",
        "[9] C. Bisaillon, 'Fake and Real News Dataset,' Kaggle, 2020. [Online]. Available: https://www.kaggle.com/datasets/clmentbisaillon/fake-and-real-news-dataset. Accessed: Oct. 16, 2026.",
        "[10] 'PolitiFact: Fact-Checking Journalism,' [Online]. Available: https://www.politifact.com. Accessed: Oct. 16, 2026."
    ]

    for ref in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)

    doc.save(DOCX_OUT)
    print(f"[SUCCESS] Saved DOCX: {DOCX_OUT}")


def convert_docx_to_pdf():
    print("[INFO] Converting DOCX to PDF via Microsoft Word COM...")
    cmd = [
        r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe",
        "-ExecutionPolicy", "Bypass",
        "-File", os.path.join(WORKSPACE_DIR, "convert_to_pdf.ps1"),
        "-InputPath", DOCX_OUT,
        "-OutputPath", PDF_OUT
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    print("PowerShell Output:", res.stdout)
    if res.stderr:
        print("PowerShell Error:", res.stderr)
    
    if os.path.exists(PDF_OUT) and os.path.getsize(PDF_OUT) > 0:
        print(f"[SUCCESS] Generated PDF: {PDF_OUT} ({os.path.getsize(PDF_OUT)} bytes)")
    else:
        print("[ERROR] PDF generation failed!")


if __name__ == "__main__":
    generate_document()
    convert_docx_to_pdf()
